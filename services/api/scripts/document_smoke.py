"""Offline renderer/parser acceptance: synthetic materials only; never calls an LLM."""
import argparse
import json
from pathlib import Path
from uuid import uuid4

from docx import Document
from PIL import Image
from archflow_api.document_worker import parse_material, render_document


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    parser.add_argument("--pages", type=int, default=5)
    parser.add_argument("--chapters", type=int, default=3)
    args = parser.parse_args()
    root = args.directory
    root.mkdir(parents=True, exist_ok=True)
    brief = Document()
    brief.add_heading("测试冷链园区任务书（虚构验收资料）", 0)
    brief.add_paragraph("项目位于测试市，用地面积 100 亩；本期建筑面积 36000 平方米。甲方希望以智慧冷链、高效物流、低碳园区作为汇报重点。")
    brief.add_paragraph("参考案例中的名称和指标不可作为本项目事实。未提供总平面、物流流线和工程计算结果，不得编造技术图纸。方案仅供设计师审阅。")
    brief.save(root / "brief.docx")
    image_dir = root / "image"
    image_dir.mkdir(exist_ok=True)
    image_id = str(uuid4())
    Image.new("RGB", (1600, 900), (52, 94, 105)).save(image_dir / "original.jpg")
    (image_dir / "metadata.json").write_text(json.dumps({"id": image_id, "name": "验收用色卡（非建筑效果图）", "role": "image"}), "utf-8")
    source = parse_material(image_dir)  # Real Pillow parsing; source catalog is identical to production.
    catalog = json.loads((image_dir / "index.json").read_text("utf-8")) | {"directory": str(image_dir)}
    units = [{"unit_index": index, "title": "冷链产业园概念方案" if index == 1 else f"设计目标与实施思路 {index}",
              "body": "智慧冷链 · 高效物流 · 低碳园区\n用地 100 亩，本期建筑面积 36000 平方米。\n概念审阅资料，工程条件需进一步确认。",
              "layout": "cover" if index == 1 else "text", "image_id": catalog["assets"][0]["id"] if index == 1 else None,
              "table": [], "evidence": ["user-brief"], "missing_facts": []} for index in range(1, args.pages+1)]
    for module, contents in [("concept", units), ("bid", [{**units[0], "unit_index": index, "title": f"技术响应与实施安排 {index}", "image_id": None, "body": "\n".join(["本段为虚构验收内容，用于验证长文档的中文字体、分页及可编辑性。实施计划应结合任务书确定责任人、检查节点和反馈机制，不得将参考项目条件作为本项目已确认事实。" * 3] * 5)} for index in range(1, args.chapters+1)])]:
        folder = root / module
        folder.mkdir(exist_ok=True)
        (folder / "input.json").write_text(json.dumps({"module": module, "title": "ArchFlow 合成验收", "units": contents, "sources": [catalog]}, ensure_ascii=False), "utf-8")
        result = render_document(folder)
        print(json.dumps({"module": module, **result}, ensure_ascii=False), flush=True)
    print(json.dumps({"image_parse": source}), flush=True)


if __name__ == "__main__":
    main()
