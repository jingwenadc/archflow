"""Single bounded CPU worker for parsing and immutable Office/PDF exports."""
import json
import os
import re
import signal
import subprocess
import tempfile
import time
import threading
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image
from pypdf import PdfReader

from .config import load_settings
from .materials import DocumentRepository


def command(args: list[str], timeout: int = 300):
    subprocess.run(args, check=True, timeout=timeout, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def office_pdf(source: Path, folder: Path, profile: Path) -> Path:
    command(["/usr/bin/libreoffice", f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(folder), str(source)])
    output = folder / f"{source.stem}.pdf"
    if not output.is_file():
        raise ValueError("Office 文件转换失败，可能损坏或受密码保护。")
    return output


def parse_material(folder: Path) -> dict:
    metadata = json.loads((folder / "metadata.json").read_text("utf-8"))
    source = next(folder.glob("original.*"))
    parsed = folder / "parsed"
    parsed.mkdir(exist_ok=True)
    pages, assets, theme = [], [], {}
    file_id = metadata["id"]
    extension = source.suffix.lower()
    if extension in {".pptx", ".docx", ".xlsx"}:
        with zipfile.ZipFile(source) as archive:
            if len(archive.infolist()) > 10000 or sum(info.file_size for info in archive.infolist()) > 300_000_000:
                raise ValueError("文件解压体积超过限制，请拆分后重试。")
    if extension in {".png", ".jpg", ".jpeg", ".webp"}:
        with Image.open(source) as image:
            if image.width * image.height > 50_000_000:
                raise ValueError("图片超过 5000 万像素，请缩小后重试。")
            image.thumbnail((1800, 1800))
            image.convert("RGB").save(parsed / "image.jpg", quality=88)
        assets.append({"id": f"{file_id}:image", "file": "parsed/image.jpg", "caption": metadata["name"]})
        pages.append({"id": f"{file_id}:p1", "page": 1, "text": "项目图片，请调用 view_image 理解内容。", "image_id": assets[0]["id"]})
    else:
        with tempfile.TemporaryDirectory(prefix="archflow-parse-") as temporary:
            temp = Path(temporary)
            pdf = source if extension == ".pdf" else office_pdf(source, temp, temp / "profile")
            reader = PdfReader(pdf)
            if reader.is_encrypted or len(reader.pages) > 500:
                raise ValueError("文件加密或超过 500 页，请解密或拆分后上传。")
            command(["/usr/bin/pdftoppm", "-jpeg", "-scale-to", "1400", str(pdf), str(parsed / "page")], timeout=600)
            images = sorted(parsed.glob("page-*.jpg"), key=lambda p: int(p.stem.split("-")[-1]))
            for index, page in enumerate(reader.pages, 1):
                text = (page.extract_text() or "").strip()
                image = images[index-1]
                if len(text) < 30:
                    try:
                        text = subprocess.check_output(["/usr/bin/tesseract", str(image), "stdout", "-l", "chi_sim+eng"], timeout=45, stderr=subprocess.DEVNULL).decode("utf-8").strip()
                    except (subprocess.SubprocessError, OSError):
                        text = "扫描页，请调用 view_image 读取原图。"
                asset_id = f"{file_id}:page{index}"
                assets.append({"id": asset_id, "file": image.relative_to(folder).as_posix(), "caption": f"{metadata['name']} 第 {index} 页"})
                pages.append({"id": f"{file_id}:p{index}", "page": index, "text": text[:30000], "image_id": asset_id})
            if extension == ".pptx":
                with zipfile.ZipFile(source) as archive:
                    if sum(info.file_size for info in archive.infolist()) > 300_000_000:
                        raise ValueError("演示文稿解压体积超过限制。")
                    for name in archive.namelist():
                        if name.startswith("ppt/media/") and Path(name).suffix.lower() in {".png", ".jpg", ".jpeg"} and len(assets) < 600:
                            target = parsed / f"asset-{len(assets)}.jpg"
                            import io
                            with Image.open(io.BytesIO(archive.read(name))) as image:
                                image.thumbnail((1800, 1800))
                                image.convert("RGB").save(target, quality=88)
                            assets.append({"id": f"{file_id}:asset{len(assets)}", "file": target.relative_to(folder).as_posix(), "caption": f"{metadata['name']} 原图 {Path(name).name}"})
                    if "ppt/theme/theme1.xml" in archive.namelist():
                        root = ET.fromstring(archive.read("ppt/theme/theme1.xml"))
                        color = root.find(".//{*}accent1/{*}srgbClr")
                        if color is not None and re.fullmatch("[A-Fa-f0-9]{6}", color.get("val", "")):
                            theme["accent"] = color.get("val")
    index = {"file_id": file_id, "name": metadata["name"], "role": metadata.get("role", "source"), "page_count": len(pages), "pages": pages, "assets": assets, "theme": theme}
    temporary_index = folder / "index.pending.json"
    temporary_index.write_text(json.dumps(index, ensure_ascii=False), "utf-8")
    temporary_index.replace(folder / "index.json")
    return {"page_count": len(pages)}


def render_document(folder: Path) -> dict:
    request = json.loads((folder / "input.json").read_text("utf-8"))
    units = request["units"]
    if not units:
        raise ValueError("尚无内容可导出。")
    with tempfile.TemporaryDirectory(prefix="archflow-export-") as temporary:
        temp = Path(temporary)
        extension = "pptx" if request["module"] == "concept" else "docx"
        output = temp / f"archflow.{extension}"
        if extension == "pptx":
            rendered = subprocess.run(["/usr/local/bin/node", "/app/documents/render-slides.mjs", str(folder / "input.json"), str(output)], capture_output=True, text=True, timeout=300)
            if rendered.returncode:
                issues = re.findall(r"第 \d+ 页[^\n]+", rendered.stderr)
                raise ValueError("；".join(issues[:10]) if issues else "演示文稿排版失败，请检查内容密度后重试。")
        else:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            document = Document()
            section = document.sections[0]
            section.page_width, section.page_height = Inches(8.5), Inches(11)
            section.top_margin = section.bottom_margin = Inches(.8)
            footer = section.footer.paragraphs[0]
            footer.alignment = 2
            page_field = OxmlElement("w:fldSimple")
            page_field.set(qn("w:instr"), "PAGE")
            footer._p.append(page_field)
            for style in ["Normal", "Title", "Heading 1", "Heading 2"]:
                font = document.styles[style].font
                font.name, font.color.rgb = "Noto Sans CJK SC", RGBColor(0, 0, 0)
                font.size = Pt(12 if style == "Normal" else 20)
                document.styles[style].element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
            for unit in units:
                document.add_heading(unit["title"], level=1)
                for paragraph in unit["body"].split("\n"):
                    if paragraph.strip():
                        document.add_paragraph(paragraph.replace("**", ""))
                if unit.get("table"):
                    table = document.add_table(rows=0, cols=max(len(row) for row in unit["table"]))
                    table.style = "Table Grid"
                    for row_index, row in enumerate(unit["table"]):
                        cells = table.add_row().cells
                        for col, value in enumerate(row):
                            cells[col].text = value
                        if row_index == 0:
                            repeat = OxmlElement("w:tblHeader")
                            table.rows[0]._tr.get_or_add_trPr().append(repeat)
                image_id = unit.get("image_id")
                for source in request["sources"]:
                    for asset in source["assets"]:
                        if asset["id"] == image_id:
                            document.add_picture(str(Path(source["directory"]) / asset["file"]), width=Inches(6))
            document.save(output)
        pdf = office_pdf(output, temp, temp / "profile")
        page_count = len(PdfReader(pdf).pages)
        if extension == "pptx" and page_count != len(units):
            raise ValueError("渲染页数与提纲不一致。")
        command(["/usr/bin/pdftoppm", "-jpeg", "-scale-to", "1100" if request.get("purpose") == "agent-preview" else "1400", str(pdf), str(temp / "preview")], timeout=600)
        import shutil
        shutil.copy2(output, folder / output.name)
        shutil.copy2(pdf, folder / "archflow.pdf")
        for index, image in enumerate(sorted(temp.glob("preview-*.jpg"), key=lambda p: int(p.stem.split("-")[-1])), 1):
            shutil.copy2(image, folder / f"page-{index}.jpg")
    return {"page_count": page_count, "format": extension, "font": "Noto Sans CJK SC", "quality": "rendered-review", "missing_facts": sum(len(unit["missing_facts"]) for unit in units)}


def main():
    settings = load_settings()
    documents = DocumentRepository(settings.database_path)
    active_work = {"id": None}
    heartbeat_stop = threading.Event()
    def heartbeat():
        while not heartbeat_stop.is_set():
            try:
                if active_work["id"]:
                    documents.renew(active_work["id"])
                Path("/tmp/archflow-documents-health").touch()
            except OSError:
                pass
            heartbeat_stop.wait(15)
    threading.Thread(target=heartbeat, daemon=True).start()
    stopping = False
    def stop(*_):
        nonlocal stopping
        stopping = True
    signal.signal(signal.SIGTERM, stop)
    signal.signal(signal.SIGINT, stop)
    Path(os.environ.get("HOME", "/tmp/archflow-documents")).mkdir(parents=True, exist_ok=True)
    while not stopping:
        work = documents.claim()
        if work:
            active_work["id"] = work["id"]
            try:
                result = parse_material(Path(work["path"])) if work["kind"] == "material" else render_document(Path(work["path"]))
                documents.finish(work["id"], result)
            except Exception as error:
                # Never persist document text, provider responses, or keys in error logs.
                message = str(error) if isinstance(error, ValueError) else "文件解析或渲染失败。请检查文件格式、内容密度或重试。"
                documents.finish(work["id"], error=message[:1000])
            finally:
                active_work["id"] = None
        else:
            time.sleep(1)


if __name__ == "__main__":
    main()
